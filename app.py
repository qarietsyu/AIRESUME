"""AI Resume ATS Checker - Streamlit + Google Gemini Flash."""

import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-2.5-flash"
MAX_RESUME_CHARS = 20000
MIN_RESUME_CHARS = 200


# --------------------------------------------------------------------------
# Resume text extraction
# --------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(BytesIO(file_bytes))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep content inside tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    elif name.endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# --------------------------------------------------------------------------
# Gemini
# --------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description.strip()}\n"
        if job_description.strip()
        else "JOB DESCRIPTION: not provided. Judge against general ATS best practices."
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.
Evaluate the resume below{" against the job description" if job_description.strip() else ""}.
Be honest and strict. Do not inflate scores. Never invent facts that are not in the resume.

Return ONLY a JSON object with exactly this structure:
{{
  "ats_score": <integer 0-100>,
  "score_breakdown": {{
    "keywords_match": <integer 0-100>,
    "formatting_and_structure": <integer 0-100>,
    "content_quality": <integer 0-100>,
    "readability": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", "..."],
  "weaknesses": ["<string>", "..."],
  "missing_keywords": ["<string>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "section": "<resume section>", "suggestion": "<specific, actionable fix>"}}
  ],
  "rewrite_examples": [
    {{"original": "<weak line copied from the resume>", "improved": "<stronger version using action verb and measurable impact, without inventing numbers>"}}
  ]
}}

Scoring guide: keywords_match (relevant skills/keywords), formatting_and_structure (standard section headings, contact info, parseable layout),
content_quality (quantified achievements, action verbs, relevance), readability (concise, consistent, no typos).
ats_score is the overall score. Give 3-6 items each for strengths, weaknesses and improvements, up to 10 missing keywords and 2-4 rewrite examples.

{jd_block}
RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating code fences or stray text."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI response was not valid JSON. Please try again.")


def _clamp_score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def normalize_result(data: dict) -> dict:
    """Make sure every field exists with the right type so the UI never crashes."""
    breakdown = data.get("score_breakdown") or {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")),
                    "section": str(item.get("section", "General")),
                    "suggestion": str(item.get("suggestion", "")),
                }
            )
        else:
            improvements.append({"priority": "Medium", "section": "General", "suggestion": str(item)})
    rewrites = [
        {"original": str(r.get("original", "")), "improved": str(r.get("improved", ""))}
        for r in (data.get("rewrite_examples") or [])
        if isinstance(r, dict)
    ]
    as_list = lambda v: [str(x) for x in v] if isinstance(v, list) else []
    return {
        "ats_score": _clamp_score(data.get("ats_score")),
        "score_breakdown": {
            "Keywords match": _clamp_score(breakdown.get("keywords_match")),
            "Formatting & structure": _clamp_score(breakdown.get("formatting_and_structure")),
            "Content quality": _clamp_score(breakdown.get("content_quality")),
            "Readability": _clamp_score(breakdown.get("readability")),
        },
        "summary": str(data.get("summary", "")),
        "strengths": as_list(data.get("strengths")),
        "weaknesses": as_list(data.get("weaknesses")),
        "missing_keywords": as_list(data.get("missing_keywords")),
        "improvements": improvements,
        "rewrite_examples": rewrites,
    }


def get_api_key(sidebar_key: str) -> str:
    if sidebar_key.strip():
        return sidebar_key.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 60:
        return "Good, with room to improve"
    if score >= 40:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["ats_score"]
    st.subheader("Your ATS score")
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Overall", f"{score}/100")
        st.caption(score_label(score))
    with col2:
        st.progress(score / 100)
        st.write(result["summary"])

    st.markdown("#### Score breakdown")
    cols = st.columns(len(result["score_breakdown"]))
    for col, (label, value) in zip(cols, result["score_breakdown"].items()):
        col.metric(label, f"{value}/100")
        col.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.markdown("#### Strengths")
        for item in result["strengths"] or ["None identified."]:
            st.markdown(f"- {item}")
    with right:
        st.markdown("#### Weaknesses")
        for item in result["weaknesses"] or ["None identified."]:
            st.markdown(f"- {item}")

    if result["missing_keywords"]:
        st.markdown("#### Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.markdown("#### Suggested improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    for item in sorted(result["improvements"], key=lambda i: order.get(i["priority"].lower(), 3)):
        st.markdown(f"**[{item['priority']}] {item['section']}**: {item['suggestion']}")

    if result["rewrite_examples"]:
        st.markdown("#### Example rewrites")
        for ex in result["rewrite_examples"]:
            st.markdown(f"- **Before:** {ex['original']}\n\n  **After:** {ex['improved']}")

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.write("Upload your resume to get an ATS score and specific suggestions to improve it.")

    with st.sidebar:
        st.header("Settings")
        sidebar_key = st.text_input(
            "Gemini API key",
            type="password",
            help="Optional if GEMINI_API_KEY is set in Streamlit secrets.",
        )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.caption("Get a free key at aistudio.google.com.")

    uploaded = st.file_uploader("Upload resume (PDF, DOCX or TXT)", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, improves keyword matching)", height=150
    )

    if st.button("Analyze resume", type="primary"):
        if uploaded is None:
            st.warning("Please upload a resume first.")
            return
        api_key = get_api_key(sidebar_key)
        if not api_key:
            st.error("No Gemini API key found. Enter it in the sidebar or add it to Streamlit secrets.")
            return

        try:
            resume_text = extract_resume_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If this is a scanned or image-based PDF, "
                "an ATS would not be able to read it either. Export a text-based PDF or DOCX."
            )
            return

        with st.spinner("Analyzing your resume..."):
            try:
                st.session_state["result"] = analyze_resume(
                    api_key, model.strip() or DEFAULT_MODEL, resume_text, job_description
                )
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")
                return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
